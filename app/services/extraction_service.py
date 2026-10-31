import logging
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader

logger = logging.getLogger(__name__)
_TERMINATORS = (".", "!", "?", "…", ":", ";", "»", '"', ")")

class ExtractionError(Exception):
    """Файл не вдалося прочитати. user_message - текст для користувача."""

    def __init__(self, user_message: str) -> None:
        super().__init__(user_message)
        self.user_message = user_message


def extract_text(path: str) -> str:
    suffix = Path(path).suffix.lower()
    try:
        if suffix == ".pdf":
            text = _extract_pdf(path)
        elif suffix == ".docx":
            text = _extract_docx(path)
        elif suffix == ".txt":
            text = _extract_txt(path)
        else:
            raise ExtractionError("Цей формат я не читаю.")
    except ExtractionError:
        raise
    except Exception as exc:
        logger.warning("Не вдалося прочитати %s", path, exc_info=True)
        raise ExtractionError(
            "Не вдалося прочитати фалй можливо він пошкоджений."
            "Надішли будь ласка, інший."
        ) from exc

    text = text.replace("\x00", "").strip()
    if not text:
        if suffix == ".pdf":
            raise ExtractionError(
                "У цьому PDF немає текстового шару — схоже, це скан. "
                "Розпізнавання тексту із зображень я поки не підтримую."
            )
        raise ExtractionError("Файл порожній — я не знайшла в ньому тексту.")
    return text

def _continues_previous(prev: str, line: str) -> bool:
    if not prev or not line or prev.endswith(_TERMINATORS):
        return False
    first = line[0]
    return first.islower() or first.isdigit()


def unwrap_lines(text: str) -> str:
    out: list[str] = []
    for line in text.splitlines():
        stripped = line.strip()
        if out and _continues_previous(out[-1], stripped):
            out[-1] = out[-1] + " " + stripped
        else:
            out.append(stripped)
    return "\n".join(out)


def _extract_pdf(path: str) -> str:
    reader = PdfReader(path)
    if reader.is_encrypted:
        raise ExtractionError("PDF захищений паролем, тому я не можу його прочитати.")
    text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
    return unwrap_lines(text)


def _extract_docx(path: str) -> str:
    document = DocxDocument(path)
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _extract_txt(path: str) -> str:
    raw = Path(path).read_bytes()
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw.decode("cp1251", errors="replace")
