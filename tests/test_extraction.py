import pytest
from docx import Document
from pypdf import PdfWriter

from app.services.extraction_service import ExtractionError, extract_text, unwrap_lines


def test_txt_utf8(tmp_path):
    f = tmp_path / "a.txt"
    f.write_text("Внесок під дубом — 4 жолуді", encoding="utf-8")
    assert extract_text(str(f)) == "Внесок під дубом — 4 жолуді"


def test_txt_cp1251(tmp_path):
    f = tmp_path / "a.txt"
    f.write_bytes("Крикни тричі".encode("cp1251"))
    assert extract_text(str(f)) == "Крикни тричі"


def test_docx_paragraphs_and_tables(tmp_path):
    f = tmp_path / "a.docx"
    doc = Document()
    doc.add_paragraph("Загальна порада")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Дуб"
    table.rows[0].cells[1].text = "4 жолуді"
    doc.save(f)

    text = extract_text(str(f))

    assert "Загальна порада" in text
    assert "Дуб | 4 жолуді" in text


def test_blank_pdf_is_reported_as_scan(tmp_path):
    f = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with open(f, "wb") as fh:
        writer.write(fh)

    with pytest.raises(ExtractionError) as exc:
        extract_text(str(f))
    assert "скан" in exc.value.user_message


def test_corrupted_pdf_gives_friendly_error(tmp_path):
    f = tmp_path / "bad.pdf"
    f.write_bytes(b"this is not a pdf")
    with pytest.raises(ExtractionError) as exc:
        extract_text(str(f))
    assert "пошкоджений" in exc.value.user_message


def test_empty_txt_is_rejected(tmp_path):
    f = tmp_path / "e.txt"
    f.write_text("   \n", encoding="utf-8")
    with pytest.raises(ExtractionError):
        extract_text(str(f))


def test_unwrap_lines_joins_broken_sentence_but_keeps_list_items():
    text = "Гриби з'являються у хвойному\nлісі на другий день.\n• Перша порада\n• Друга порада\nЗаголовок\nТекст"
    assert unwrap_lines(text).splitlines() == [
        "Гриби з'являються у хвойному лісі на другий день.",
        "• Перша порада",
        "• Друга порада",
        "Заголовок",
        "Текст",
    ]
